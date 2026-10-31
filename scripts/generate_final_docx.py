import os
import sys

sys.path.insert(0, os.path.abspath("."))

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from scripts.docx_styles import (
    NAVY_BLUE, DEEP_BLUE, set_cell_background, 
    set_cell_margins, add_explanation_box
)
from scripts.build_part1_part2 import add_part_1, add_part_2
from scripts.build_part3_part4_part5 import add_part_3, add_part_4, add_part_5

def build_master_docx(output_path):
    print(f"Building master Word document at: {output_path}")
    doc = docx.Document()
    
    # 1. Cấu hình lề trang 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = False
        
    # 2. Trang bìa / Header đầu tài liệu
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_before = Pt(0)
    p_univ.paragraph_format.space_after = Pt(2)
    run_u = p_univ.add_run("TRƯỜNG ĐẠI HỌC BÁCH KHOA - ĐẠI HỌC ĐÀ NẴNG\nKHOA CÔNG NGHỆ THÔNG TIN")
    run_u.font.name = "Calibri"
    run_u.font.size = Pt(11)
    run_u.font.bold = True
    run_u.font.color.rgb = RGBColor(100, 116, 139)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run("KỊCH BẢN THUYẾT TRÌNH VÀ THAO TÁC BẢO VỆ TIỂU LUẬN\nQUẢN TRỊ CƠ SỞ DỮ LIỆU")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(17)
    run_t.font.bold = True
    run_t.font.color.rgb = NAVY_BLUE
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(12)
    run_s = p_sub.add_run("ĐỀ TÀI: HỆ THỐNG QUẢN LÝ PHÒNG GAME / NET (NETCAFE MANAGEMENT SYSTEM)\n"
                          "Nền tảng: PostgreSQL 16 (Neon Cloud) | Python 3.12 ETL | pg_trgm | LISTEN/NOTIFY | 5 Module AI")
    run_s.font.name = "Calibri"
    run_s.font.size = Pt(10.5)
    run_s.font.italic = True
    run_s.font.color.rgb = DEEP_BLUE
    
    # 3. Khung hướng dẫn chuẩn bị môi trường
    add_explanation_box(
        doc,
        "LƯU Ý QUAN TRỌNG TRƯỚC KHI THAO TÁC TRÊN MÁY",
        "1. Chuẩn bị 2 cửa sổ song song trên màn hình:\n"
        "   - Cửa sổ NỬA TRÁI: Trình duyệt mở Neon SQL Editor (hoặc DBeaver kết nối database netcafe_db).\n"
        "   - Cửa sổ NỬA PHẢI: Terminal (PowerShell) trỏ tại thư mục gốc dự án C:\\Users\\Administrator\\Documents\\QTDL-NetCafe.\n"
        "2. Thứ tự thực hiện: Chạy tuần tự từ Bước 1.1 đến Bước 4.4. Người thuyết trình đọc lời dẫn giải trước, "
        "sau đó copy lệnh tương ứng dán vào cửa sổ thao tác để chứng minh trên CSDL thực tế."
    )
    
    # 4. Thêm từng phần có cấu trúc mạch lạc
    add_part_1(doc)
    add_part_2(doc)
    add_part_3(doc)
    add_part_4(doc)
    add_part_5(doc)
    
    # 5. Lưu tài liệu
    for path in [output_path, output_path.replace(".docx", "_Pending.docx"), output_path.replace(".docx", "_Latest.docx")]:
        try:
            doc.save(path)
            print(f"Successfully generated: {path}")
            break
        except PermissionError:
            continue

if __name__ == "__main__":
    out_file = os.path.join("docs", "KichBan_ThuyetTrinh_BaoVe_QTDL_MoiNhat.docx")
    build_master_docx(out_file)
